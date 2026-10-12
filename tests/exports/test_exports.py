"""
TDD: 导出模块
"""
import pytest
import os


class TestMarkdownExport:
    def test_md_ai_analysis(self, empty_duckdb_conn):
        from exports.markdown import ai_analysis_to_md
        empty_duckdb_conn.execute("""
            INSERT INTO ai_analysis
                (analysis_id, stock_code, analysis_type, prompt, response, model_used)
            VALUES ('a1', '000001', 'financial', '你是一个分析师',
                    '# 标题\n正文内容\n## 子标题\n- 项目 1\n- 项目 2', 'mock')
        """)
        md = ai_analysis_to_md(empty_duckdb_conn, "a1")
        assert "# 标题" in md
        assert "## 子标题" in md
        assert "000001" in md
        assert "mock" in md

    def test_md_note(self, empty_duckdb_conn):
        from exports.markdown import note_to_md
        empty_duckdb_conn.execute("""
            INSERT INTO note (note_id, stock_code, title, content_md, synced_to_vector)
            VALUES ('n1', '000001', '我的观察', '内容...', false)
        """)
        md = note_to_md(empty_duckdb_conn, "n1")
        assert "我的观察" in md
        assert "内容..." in md

    def test_md_watchlist(self, empty_duckdb_conn):
        from exports.markdown import watchlist_to_md
        empty_duckdb_conn.execute("""
            INSERT INTO stock_basic (stock_code, short_name, exchange)
            VALUES ('000001', '平安银行', 'SZ'), ('000002', '万科A', 'SZ')
        """)
        empty_duckdb_conn.execute("""
            INSERT INTO watchlist_group (group_name, description, sort_order)
            VALUES ('持仓池', '我的持仓', 1)
        """)
        empty_duckdb_conn.execute("""
            INSERT INTO watchlist (group_name, stock_code, note) VALUES
            ('持仓池', '000001', '低估值'), ('持仓池', '000002', '观察')
        """)
        md = watchlist_to_md(empty_duckdb_conn, "持仓池")
        assert "持仓池" in md
        assert "平安银行" in md
        assert "万科A" in md
        assert "低估值" in md

    def test_md_file_write(self, empty_duckdb_conn, tmp_path):
        from exports.markdown import ai_analysis_to_md_file
        empty_duckdb_conn.execute("""
            INSERT INTO ai_analysis
                (analysis_id, stock_code, analysis_type, response, model_used)
            VALUES ('a1', '000001', 'financial', '# 标题\n正文', 'mock')
        """)
        path = ai_analysis_to_md_file(empty_duckdb_conn, "a1", str(tmp_path))
        assert os.path.exists(path)
        assert path.endswith(".md")
        with open(path, encoding="utf-8") as f:
            content = f.read()
        assert "# 标题" in content

    def test_md_not_found_raises(self, empty_duckdb_conn):
        from exports.markdown import ai_analysis_to_md
        with pytest.raises(KeyError):
            ai_analysis_to_md(empty_duckdb_conn, "nonexistent")


class TestPDFExport:
    def test_pdf_from_markdown(self, tmp_path):
        from exports.pdf import md_to_pdf
        md = "# 标题\n## 子标题\n正文内容。\n\n- 项目 1\n- 项目 2"
        path = md_to_pdf(md, str(tmp_path), "test.pdf")
        assert os.path.exists(path)
        assert path.endswith(".pdf")
        assert os.path.getsize(path) > 500

    def test_pdf_chinese_supported(self, tmp_path):
        from exports.pdf import md_to_pdf
        md = "# 财报分析\n\n 营业收入 100 亿元，同比增长 20%。"
        path = md_to_pdf(md, str(tmp_path), "chinese.pdf")
        assert os.path.exists(path)
        assert os.path.getsize(path) > 500

    def test_pdf_from_ai_analysis(self, empty_duckdb_conn, tmp_path):
        from exports.pdf import ai_analysis_to_pdf
        empty_duckdb_conn.execute("""
            INSERT INTO ai_analysis
                (analysis_id, stock_code, analysis_type, response, model_used)
            VALUES ('a1', '000001', 'financial', '# 平安银行\n\n 营收 100 亿。', 'mock')
        """)
        path = ai_analysis_to_pdf(empty_duckdb_conn, "a1", str(tmp_path))
        assert os.path.exists(path)
        assert path.endswith(".pdf")


class TestExcelExport:
    def test_excel_watchlist(self, empty_duckdb_conn, tmp_path):
        from exports.excel import watchlist_to_excel
        empty_duckdb_conn.execute("""
            INSERT INTO stock_basic (stock_code, short_name, exchange) VALUES
            ('000001', '平安银行', 'SZ'), ('000002', '万科A', 'SZ')
        """)
        empty_duckdb_conn.execute("""
            INSERT INTO watchlist_group (group_name, description, sort_order)
            VALUES ('持仓池', '我的持仓', 1)
        """)
        empty_duckdb_conn.execute("""
            INSERT INTO watchlist (group_name, stock_code, note) VALUES
            ('持仓池', '000001', '低估值'), ('持仓池', '000002', '观察')
        """)
        path = watchlist_to_excel(empty_duckdb_conn, "持仓池", str(tmp_path))
        assert os.path.exists(path)
        assert path.endswith(".xlsx")
        assert os.path.getsize(path) > 1000

    def test_excel_portfolio(self, empty_duckdb_conn, tmp_path):
        from exports.excel import portfolio_to_excel
        empty_duckdb_conn.execute("""
            INSERT INTO stock_basic (stock_code, short_name, exchange) VALUES
            ('000001', '平安银行', 'SZ')
        """)
        empty_duckdb_conn.execute("""
            INSERT INTO position (position_id, stock_code, cost_price, quantity, cost_total, open_date, note, status)
            VALUES ('p1', '000001', 10.5, 100, 1050.0, '2024-01-01', '初始建仓', 'open')
        """)
        empty_duckdb_conn.execute("""
            INSERT INTO trade (trade_id, stock_code, direction, price, quantity, amount, trade_date, logic_text)
            VALUES ('t1', '000001', 'buy', 10.5, 100, 1050.0, '2024-01-01', '建仓')
        """)
        path = portfolio_to_excel(empty_duckdb_conn, str(tmp_path))
        assert os.path.exists(path)
        assert path.endswith(".xlsx")
        assert os.path.getsize(path) > 1000

    def test_excel_contains_data(self, empty_duckdb_conn, tmp_path):
        from openpyxl import load_workbook
        from exports.excel import watchlist_to_excel
        empty_duckdb_conn.execute("""
            INSERT INTO stock_basic (stock_code, short_name, exchange) VALUES
            ('000001', '平安银行', ' ')
        """)
        empty_duckdb_conn.execute("""
            INSERT INTO watchlist_group (group_name, sort_order) VALUES ('g1', 1)
        """)
        empty_duckdb_conn.execute("""
            INSERT INTO watchlist (group_name, stock_code, note) VALUES ('g1', '000001', 'test')
        """)
        path = watchlist_to_excel(empty_duckdb_conn, "g1", str(tmp_path))
        wb = load_workbook(path)
        ws = wb.active
        all_text = []
        for row in ws.iter_rows(values_only=True):
            for c in row:
                if c is not None:
                    all_text.append(str(c))
        all_joined = "\n".join(all_text)
        # 表头 + 数据 至少有一个 "代码" / "平安银行"
        assert "代码" in all_joined or "平安银行" in all_joined
        assert "test" in all_joined
