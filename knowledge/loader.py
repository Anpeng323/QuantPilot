"""
文档加载：支持 PDF / Word / Markdown / TXT / HTML。
返回 Document 列表（每个文件可能产出多个 page/segment）。
"""
from __future__ import annotations
import re
from pathlib import Path
from dataclasses import dataclass, field
from typing import List


@dataclass
class DocSegment:
    text: str
    page: int = 0       # PDF 页码（其他类型为 0）
    section: str = ""   # 小节标题（尽力提取）
    metadata: dict = field(default_factory=dict)


def load_pdf(path: str) -> List[DocSegment]:
    import pymupdf
    segs = []
    with pymupdf.open(path) as doc:
        for i, page in enumerate(doc, start=1):
            text = _normalize(page.get_text("text") or "")
            if text.strip():
                segs.append(DocSegment(text=text, page=i))
    # 有文本层则直接返回；否则走 OCR 兜底（扫描件 PDF）
    if segs:
        return segs
    print("[loader] PDF 无文本层，尝试 OCR 识别（扫描件）...")
    return ocr_pdf(path)


def ocr_pdf(path: str, dpi: int = 200) -> List[DocSegment]:
    """扫描件 PDF 兜底：逐页渲染为图像，用 rapidocr 识别文字。"""
    try:
        from rapidocr_onnxruntime import RapidOCR
    except Exception as e:
        print(f"[loader] OCR 不可用（未安装 rapidocr-onnxruntime）：{e}")
        return []
    import numpy as np
    import pymupdf
    try:
        engine = RapidOCR()
    except Exception as e:
        print(f"[loader] OCR 引擎初始化失败：{e}")
        return []
    segs = []
    with pymupdf.open(path) as doc:
        for i, page in enumerate(doc, start=1):
            pix = page.get_pixmap(dpi=dpi)
            arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
            if pix.n == 4:
                arr = arr[..., :3]
            arr_bgr = arr[..., ::-1]  # RGB → BGR（OpenCV 习惯）
            try:
                result, _ = engine(arr_bgr)
            except Exception as e:
                print(f"[loader] 第 {i} 页 OCR 失败：{e}")
                result = None
            if not result:
                continue
            text = "\n".join([line[1] for line in result if line and len(line) > 1])
            text = _normalize(text)
            if text.strip():
                segs.append(DocSegment(text=text, page=i))
    return segs


def load_docx(path: str) -> List[DocSegment]:
    import docx
    d = docx.Document(path)
    segs = []
    for i, p in enumerate(d.paragraphs, start=1):
        t = _normalize(p.text)
        if t.strip():
            segs.append(DocSegment(text=t, page=0, section=p.style.name if p.style else ""))
    # 表格
    for tbl_i, tbl in enumerate(d.tables, start=1):
        rows = []
        for row in tbl.rows:
            rows.append(" | ".join(cell.text.strip() for cell in row.cells))
        t = _normalize("\n".join(rows))
        if t.strip():
            segs.append(DocSegment(text=t, page=0, section=f"Table {tbl_i}"))
    return segs


def load_md_or_txt(path: str) -> List[DocSegment]:
    text = Path(path).read_text(encoding="utf-8", errors="ignore")
    return [DocSegment(text=_normalize(text))]


def load_html(path: str) -> List[DocSegment]:
    text = Path(path).read_text(encoding="utf-8", errors="ignore")
    text = re.sub(r"<script[\s\S]*?</script>", " ", text, flags=re.I)
    text = re.sub(r"<style[\s\S]*?</style>",  " ", text, flags=re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    return [DocSegment(text=_normalize(text))]


def _normalize(text: str) -> str:
    """统一空白字符、去除控制字符"""
    if not text:
        return ""
    text = re.sub(r"[\x00-\x08\x0b-\x1f\x7f]", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def load(path: str) -> List[DocSegment]:
    """根据扩展名自动选择 loader。"""
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return load_pdf(path)
    if ext == ".docx":
        return load_docx(path)
    if ext in (".md", ".markdown", ".txt"):
        return load_md_or_txt(path)
    if ext in (".html", ".htm"):
        return load_html(path)
    raise ValueError(f"不支持的文件类型: {ext}")
