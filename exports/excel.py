"""
Excel 导出（openpyxl）。
"""
from __future__ import annotations
import os
from pathlib import Path
from typing import List, Any

import duckdb
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


_HEADER_FONT  = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
_HEADER_FILL  = PatternFill("solid", fgColor="0A2540")
_CELL_ALIGN_C = Alignment(horizontal="center", vertical="center")
_CELL_ALIGN_L = Alignment(horizontal="left",   vertical="center", wrap_text=True)
_THIN_BORDER  = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)


def _style_header(ws, headers: List[str]) -> None:
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col_idx, value=h)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.alignment = _CELL_ALIGN_C
        c.border = _THIN_BORDER


def _style_rows(ws, n_rows: int, n_cols: int, num_cols: List[int] = None) -> None:
    num_cols = num_cols or []
    for r in range(2, n_rows + 2):
        for c in range(1, n_cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = _CELL_ALIGN_C if c in num_cols else _CELL_ALIGN_L
            cell.border = _THIN_BORDER


def _autosize(ws, max_width: int = 30) -> None:
    for col in ws.columns:
        max_len = 8
        col_letter = get_column_letter(col[0].column)
        for c in col:
            v = c.value
            if v is None:
                continue
            try:
                s = str(v)
                # 中文宽度估算：1 个汉字 ≈ 2 个拉丁字符
                width = sum(2 if ord(ch) > 127 else 1 for ch in s)
                max_len = max(max_len, width)
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, max_width)


def watchlist_to_excel(conn: duckdb.DuckDBPyConnection, group_name: str, out_dir: str) -> str:
    """自选池 → xlsx"""
    Path(out_dir).mkdir(parents=True, exist_ok=True)
    grp = conn.execute(
        "SELECT description FROM watchlist_group WHERE group_name=?", [group_name]
    ).fetchone()
    rows = conn.execute("""
        SELECT w.stock_code, COALESCE(s.short_name, ''), COALESCE(w.note, ''),
               COALESCE(s.exchange, ''), w.added_at
        FROM watchlist w LEFT JOIN stock_basic s ON w.stock_code = s.stock_code
        WHERE w.group_name = ?
        ORDER BY w.added_at DESC
    """, [group_name]).fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = group_name[:30]
    headers = ["代码", "名称", "备注", "市场", "加入时间"]
    _style_header(ws, headers)
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)
    _style_rows(ws, len(rows), len(headers), num_cols=[1, 4, 5])
    _autosize(ws)
    # 在第 1 行前加一行分组描述
    if grp:
        ws.insert_rows(1)
        ws.cell(row=1, column=1, value=f"{group_name}：{grp[0] or ''}")
        ws.cell(row=1, column=1).font = Font(bold=True, size=12, color="0A2540")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))

    path = os.path.join(out_dir, f"watchlist_{group_name}.xlsx")
    wb.save(path)
    return path


def portfolio_to_excel(conn: duckdb.DuckDBPyConnection, out_dir: str) -> str:
    """持仓 + 交易 → xlsx（两个 sheet）"""
    Path(out_dir).mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # Sheet 1: 持仓
    ws = wb.active
    ws.title = "持仓"
    pos_headers = ["代码", "名称", "成本价", "数量", "成本总额", "建仓日期", "状态", "备注"]
    _style_header(ws, pos_headers)
    pos_rows = conn.execute("""
        SELECT p.stock_code, COALESCE(s.short_name, ''), p.cost_price, p.quantity,
               p.cost_total, p.open_date, p.status, COALESCE(p.note, '')
        FROM position p LEFT JOIN stock_basic s ON p.stock_code = s.stock_code
        ORDER BY p.open_date DESC
    """).fetchall()
    for r_idx, row in enumerate(pos_rows, start=2):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)
    _style_rows(ws, len(pos_rows), len(pos_headers), num_cols=[1, 3, 4, 5, 6, 7])
    _autosize(ws)

    # Sheet 2: 交易
    ws2 = wb.create_sheet("交易记录")
    trade_headers = ["日期", "代码", "名称", "方向", "价格", "数量", "金额", "买入/卖出逻辑"]
    _style_header(ws2, trade_headers)
    trade_rows = conn.execute("""
        SELECT t.trade_date, t.stock_code, COALESCE(s.short_name, ''), t.direction,
               t.price, t.quantity, t.amount, COALESCE(t.logic_text, '')
        FROM trade t LEFT JOIN stock_basic s ON t.stock_code = s.stock_code
        ORDER BY t.trade_date DESC
    """).fetchall()
    for r_idx, row in enumerate(trade_rows, start=2):
        for c_idx, val in enumerate(row, start=1):
            ws2.cell(row=r_idx, column=c_idx, value=val)
    _style_rows(ws2, len(trade_rows), len(trade_headers), num_cols=[1, 2, 4, 5, 6, 7])
    _autosize(ws2)

    path = os.path.join(out_dir, "portfolio.xlsx")
    wb.save(path)
    return path
